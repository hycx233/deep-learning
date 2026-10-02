#!/usr/bin/env python3
"""将 16:9 Beamer PDF 逐页导出成可播放的图片式 PPTX；内容仍在 LaTeX 中编辑。"""

import argparse
from pathlib import Path
import shutil
import subprocess
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-pdf", type=Path, required=True)
    parser.add_argument("--output-pptx", type=Path, required=True)
    parser.add_argument("--width-px", type=int, default=2560,
                        help="每页 PNG 宽度，默认 2560 像素")
    args = parser.parse_args()
    if not args.input_pdf.is_file():
        parser.error(f"找不到 PDF：{args.input_pdf}")
    if args.output_pptx.suffix.lower() != ".pptx":
        parser.error("输出路径应以 .pptx 结尾")
    if args.width_px < 1:
        parser.error("--width-px 必须大于零")
    renderer = shutil.which("pdftoppm")
    if renderer is None:
        parser.error("未找到 pdftoppm，请安装 Poppler 并加入 PATH；见 docs/slides/README.md")

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = args.input_pdf.stem
    deck.core_properties.subject = "Beamer 页面导出；编辑源文件后重新生成"
    with tempfile.TemporaryDirectory(prefix="beamer-pptx-") as work:
        prefix = Path(work) / "page"
        subprocess.run([renderer, "-png", "-scale-to-x", str(args.width_px),
                        "-scale-to-y", "-1", str(args.input_pdf), str(prefix)], check=True)
        pages = sorted(Path(work).glob("page-*.png"),
                       key=lambda path: int(path.stem.rsplit("-", 1)[1]))
        if not pages:
            raise ValueError(f"PDF 没有渲染出页面：{args.input_pdf}")
        for page in pages:
            with Image.open(page) as picture:
                width, height = picture.size
            if abs(width / height - 16 / 9) > 0.01:
                raise ValueError(f"{page.name} 不是 16:9，请检查 Beamer 的 aspectratio=169")
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(page), 0, 0,
                                     width=deck.slide_width, height=deck.slide_height)
        args.output_pptx.parent.mkdir(parents=True, exist_ok=True)
        deck.save(args.output_pptx)
    print(f"已导出 {len(deck.slides)} 页：{args.output_pptx}（整页图片，内容在 LaTeX 中编辑）")


if __name__ == "__main__":
    main()

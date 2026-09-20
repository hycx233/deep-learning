#!/usr/bin/env python3
"""从课程 Excel 只读提取结构和基因标注，坐标依据见 data/README.md。"""

import argparse
import csv
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


CHROM = "NC_000913.3"
CHROM_LENGTH = 4_641_652
STRUCTURE_TABLES = {"OPCID": (4, 68), "CHIN": (5, 250), "CHID": (6, 26)}
STRUCTURE_FIELDS = [
    "ID", "type", "chrom", "start", "end", "center", "length_bp",
    "source_sheet", "source_row", "source_chrom", "source_start", "source_end",
    "source_center",
]
GENE_FIELDS = [
    "gene_id", "name", "chrom", "start", "end", "strand", "length_bp",
    "source_sheet", "source_row", "source_chrom", "source_start", "source_end",
]


def table_rows(workbook, sheet_name, headers):
    """只读取所需列，忽略 Excel 尾部仅含格式的空行。"""
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(max_col=len(headers), values_only=True)
    actual_headers = next(rows)
    if tuple(headers) != actual_headers:
        raise ValueError(f"{sheet_name}: 表头应为 {headers}，实际为 {actual_headers}")
    for row_number, values in enumerate(rows, start=2):
        if all(value is None for value in values):
            continue
        if any(value is None for value in values):
            raise ValueError(f"{sheet_name} 第 {row_number} 行: 必需字段缺失")
        yield row_number, dict(zip(headers, values))


def integer_coordinate(value, location):
    if not isinstance(value, (int, float)) or not float(value).is_integer():
        raise ValueError(f"{location}: 坐标必须为整数，实际为 {value!r}")
    return int(value)


def validate_region(source_chrom, start, end, location):
    if source_chrom not in {"MG1655", CHROM}:
        raise ValueError(f"{location}: 未知染色体名 {source_chrom!r}")
    if not 0 <= start < end <= CHROM_LENGTH:
        raise ValueError(f"{location}: 区间 [{start}, {end}) 超出染色体或长度非正")


def read_structures(workbook):
    records = []
    seen_ids = set()
    seen_regions = set()
    for structure_type, (table_number, expected_count) in STRUCTURE_TABLES.items():
        sheet_name = f"Supplementary Table {table_number}"
        id_header = f"{structure_type}_ID"
        headers = [id_header, "Chr", "Start", "End"]
        if structure_type == "CHIN":
            headers = [id_header, "Chr", "Start", "Center", "End"]
        count = 0
        for row_number, row in table_rows(workbook, sheet_name, headers):
            location = f"{sheet_name} 第 {row_number} 行"
            structure_id = str(row[id_header])
            if structure_id in seen_ids:
                raise ValueError(f"{location}: 重复结构 ID {structure_id}")
            seen_ids.add(structure_id)
            # 作者按接触图坐标直接分箱；保留数值，项目内统一解释为 [start, end)。
            start = integer_coordinate(row["Start"], location)
            end = integer_coordinate(row["End"], location)
            validate_region(row["Chr"], start, end, location)
            source_center = ""
            if structure_type == "CHIN":
                source_center = integer_coordinate(row["Center"], location)
                center = source_center
            else:
                center = (start + end) // 2
            if not start <= center < end:
                raise ValueError(f"{location}: Center {center} 不在区间内")
            region_key = (structure_type, start, end)
            if region_key in seen_regions:
                raise ValueError(f"{location}: {structure_type} 存在重复区间")
            seen_regions.add(region_key)
            records.append({
                "ID": structure_id, "type": structure_type, "chrom": CHROM,
                "start": start, "end": end, "center": center,
                "length_bp": end - start, "source_sheet": sheet_name,
                "source_row": row_number, "source_chrom": row["Chr"],
                "source_start": start, "source_end": end,
                "source_center": source_center,
            })
            count += 1
        if count != expected_count:
            raise ValueError(f"{sheet_name}: 应有 {expected_count} 条，实际为 {count}")
    return records


def read_genes(workbook):
    sheet_name = "Supplementary Table 1"
    headers = ["Gene", "Chr", "start", "end", "strand"]
    records = []
    seen_ids = set()
    for row_number, row in table_rows(workbook, sheet_name, headers):
        location = f"{sheet_name} 第 {row_number} 行"
        source_start = integer_coordinate(row["start"], location)
        source_end = integer_coordinate(row["end"], location)
        # Table 1 的基因区间为 1-based 闭区间，转换时仅起点减一。
        start, end = source_start - 1, source_end
        validate_region(row["Chr"], start, end, location)
        strand = row["strand"]
        if strand not in {"+", "-"}:
            raise ValueError(f"{location}: 未知链方向 {strand!r}")
        name = str(row["Gene"])
        # 同名基因可能位于不同位置，使用源坐标和链方向区分，不按名称去重。
        gene_id = f"{name}:{source_start}-{source_end}:{strand}"
        if gene_id in seen_ids:
            raise ValueError(f"{location}: 重复基因记录 {gene_id}")
        seen_ids.add(gene_id)
        records.append({
            "gene_id": gene_id, "name": name, "chrom": CHROM,
            "start": start, "end": end, "strand": strand,
            "length_bp": end - start, "source_sheet": sheet_name,
            "source_row": row_number, "source_chrom": row["Chr"],
            "source_start": source_start, "source_end": source_end,
        })
    if not records:
        raise ValueError(f"{sheet_name}: 未读取到基因记录")
    return records


def write_csv(path, fields, records):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(records)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data/标注数据.xlsx"))
    parser.add_argument("--output-dir", type=Path, default=Path("data"))
    args = parser.parse_args()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    try:
        structures = read_structures(workbook)
        genes = read_genes(workbook)
    finally:
        workbook.close()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "structures.csv", STRUCTURE_FIELDS, structures)
    write_csv(args.output_dir / "genes.csv", GENE_FIELDS, genes)
    print(f"结构: {dict(Counter(row['type'] for row in structures))}，合计 {len(structures)}")
    names = Counter(row["name"] for row in genes)
    print(f"基因: {len(genes)} 条，{len(names)} 个不同名称（同名不同位置完整保留）")
    for structure_type in STRUCTURE_TABLES:
        example = next(row for row in structures if row["type"] == structure_type)
        print(
            f"示例 {example['ID']}: {example['chrom']}:[{example['start']}, "
            f"{example['end']}), center={example['center']}, "
            f"来源 {example['source_sheet']} 第 {example['source_row']} 行"
        )
    print(f"输出: {args.output_dir / 'structures.csv'}, {args.output_dir / 'genes.csv'}")


if __name__ == "__main__":
    main()
